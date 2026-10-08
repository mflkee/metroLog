"""Equipment-exports mixin for the equipment service (Excel import/export)."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import TYPE_CHECKING

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook

from app.models.equipment import EquipmentStatus, EquipmentType
from app.models.event import EventCategory
from app.models.user import User
from app.schemas.arshin import ArshinSearchResultRead, ArshinVriDetailRead
from app.schemas.equipment import (
    EquipmentCreateRequest,
    EquipmentSIBulkImportResultRead,
    EquipmentSIBulkImportRowRead,
)
from app.services.arshin_service import ArshinService

if TYPE_CHECKING:
    from sqlalchemy.orm import Session

    from app.models.equipment import Equipment, EquipmentFolder
    from app.repositories.equipment_repository import (
        EquipmentFolderRepository,
    )


def _create_temp_export_file_path(*, suffix: str) -> Path:
    temp_file = NamedTemporaryFile(delete=False, suffix=suffix)
    temp_file.close()
    return Path(temp_file.name)


@dataclass(slots=True)
class ParsedCertificateImportRow:
    row_number: int
    certificate_number: str
    verification_year: int | None = None


def _build_si_create_request_from_arshin(
    *,
    folder_id: int,
    object_name: str,
    status_value: EquipmentStatus,
    current_location_manual: str | None,
    result: ArshinSearchResultRead,
    detail: ArshinVriDetailRead,
) -> EquipmentCreateRequest:
    return EquipmentCreateRequest(
        folder_id=folder_id,
        group_id=None,
        object_name=object_name,
        equipment_type=EquipmentType.SI,
        name=detail.type_name or result.mit_title or "СИ из Аршина",
        modification=detail.modification or result.mi_modification,
        serial_number=detail.serial_number or result.mi_number,
        manufacture_year=detail.manufacture_year,
        status=status_value,
        current_location_manual=current_location_manual,
        si_verification={
            "vri_id": result.vri_id,
            "arshin_url": result.arshin_url,
            "org_title": result.org_title,
            "mit_number": result.mit_number,
            "mit_title": result.mit_title,
            "mit_notation": result.mit_notation,
            "mi_number": result.mi_number,
            "certificate_number": detail.certificate_number or result.result_docnum,
            "result_docnum": result.result_docnum,
            "verification_date": result.verification_date,
            "valid_date": result.valid_date,
            "raw_payload_json": result.raw_payload_json,
            "detail_payload_json": detail.raw_payload_json,
        },
    )


def _detect_certificate_column(header_row: list[str]) -> int | None:
    for index, value in enumerate(header_row):
        normalized = _normalize_header_label(value)
        if not normalized:
            continue
        if any(
            keyword in normalized for keyword in ("свид", "certificate", "документ", "document")
        ):
            return index
    return None


def _detect_certificate_layout(
    rows: list[list[str]],
) -> tuple[int | None, int | None, int | None]:
    max_header_scan_rows = min(len(rows), 20)
    for row_index in range(max_header_scan_rows):
        cert_column = _detect_certificate_column(rows[row_index])
        if cert_column is not None:
            date_column = _detect_verification_date_column(rows[row_index])
            return row_index, cert_column, date_column
    return None, None, None


def _detect_verification_date_column(header_row: list[str]) -> int | None:
    for index, value in enumerate(header_row):
        normalized = _normalize_header_label(value)
        if not normalized:
            continue
        if normalized == "дата поверки" or normalized == "verification date":
            return index
        if "дата" in normalized and "поверк" in normalized:
            return index
        if "verification" in normalized and "date" in normalized:
            return index
    return None


def _extract_certificate_rows_from_csv(file_path: Path) -> list[ParsedCertificateImportRow]:
    with file_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))
    return _extract_certificate_rows_from_matrix(rows)


def _extract_certificate_rows_from_matrix(
    rows: list[list[str]],
) -> list[ParsedCertificateImportRow]:
    if not rows:
        return []

    header_row_index, cert_column, date_column = _detect_certificate_layout(rows)
    start_index = header_row_index + 1 if header_row_index is not None else 0
    if cert_column is None:
        cert_column = 0

    extracted: list[ParsedCertificateImportRow] = []
    for index, row in enumerate(rows[start_index:], start=start_index + 1):
        if cert_column >= len(row):
            continue
        certificate_number = row[cert_column].strip()
        if not certificate_number:
            continue
        verification_year = None
        if date_column is not None and date_column < len(row):
            verification_year = _extract_year_from_cell(row[date_column])
        extracted.append(
            ParsedCertificateImportRow(
                row_number=index,
                certificate_number=certificate_number,
                verification_year=verification_year,
            )
        )
    return extracted


def _extract_certificate_rows_from_table(
    *,
    file_name: str | None,
    file_path: Path,
) -> list[ParsedCertificateImportRow]:
    lower_name = (file_name or "").lower()
    if lower_name.endswith(".csv"):
        return _extract_certificate_rows_from_csv(file_path)
    return _extract_certificate_rows_from_workbook(file_path)


def _extract_certificate_rows_from_workbook(file_path: Path) -> list[ParsedCertificateImportRow]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        matrix: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            matrix.append(["" if cell is None else str(cell).strip() for cell in row])
        return _extract_certificate_rows_from_matrix(matrix)
    finally:
        workbook.close()


def _extract_year_from_cell(value: str) -> int | None:
    match = re.search(r"(19|20)\d{2}", value)
    if not match:
        return None
    year = int(match.group(0))
    if 1900 <= year <= 2100:
        return year
    return None


def _normalize_certificate_number(value: str | None) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", value).upper()


def _normalize_header_label(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def _save_workbook_to_temp_file(workbook: Workbook) -> Path:
    temp_path = _create_temp_export_file_path(suffix=".xlsx")
    try:
        workbook.save(temp_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return temp_path


def _select_bulk_import_candidate(
    *,
    certificate_number: str,
    results: list[ArshinSearchResultRead],
) -> ArshinSearchResultRead | None:
    if not results:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Arshin did not return any records for this certificate number.",
        )

    if len(results) == 1:
        return results[0]

    normalized_input = _normalize_certificate_number(certificate_number)
    exact_matches = [
        result
        for result in results
        if _normalize_certificate_number(result.result_docnum) == normalized_input
    ]
    if len(exact_matches) == 1:
        return exact_matches[0]

    return None


class EquipmentExportsMixin:
    """Mixed into ``EquipmentService``."""

    if TYPE_CHECKING:
        session: Session
        folders: EquipmentFolderRepository

        def _get_folder(self, folder_id: int) -> EquipmentFolder: ...

        def _record_event(
            self,
            *,
            category: EventCategory,
            action: str,
            user: User | None,
            title: str,
            description: str | None = None,
            equipment: Equipment | None = None,
            equipment_id: int | None = None,
            equipment_name: str | None = None,
            equipment_modification: str | None = None,
            equipment_serial_number: str | None = None,
            folder_id: int | None = None,
            folder_name: str | None = None,
            notification_equipment_ids: list[int] | None = None,
            batch_key: str | None = None,
        ) -> None: ...

        def create_equipment(
            self,
            payload: EquipmentCreateRequest,
            *,
            current_user: User | None = None,
        ) -> Equipment: ...

        def list_equipment(
            self,
            *,
            folder_id: int | None = None,
            group_id: int | None = None,
            equipment_ids: list[int] | None = None,
            query: str | None = None,
            object_name: str | None = None,
            current_location_manual: str | None = None,
            status: EquipmentStatus | None = None,
            equipment_type: EquipmentType | None = None,
        ) -> list[Equipment]: ...

    async def import_si_from_excel(
        self,
        *,
        file_name: str | None,
        file_path: Path,
        folder_id: int,
        object_name: str,
        status_value: EquipmentStatus,
        current_location_manual: str | None,
        arshin_service: ArshinService | None = None,
        current_user: User | None = None,
    ) -> EquipmentSIBulkImportResultRead:
        self._get_folder(folder_id)

        rows = _extract_certificate_rows_from_table(
            file_name=file_name,
            file_path=file_path,
        )
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Excel file does not contain certificate numbers.",
            )

        arshin = arshin_service or ArshinService()
        result_rows: list[EquipmentSIBulkImportRowRead] = []
        seen_certificates: set[str] = set()

        for row in rows:
            row_number = row.row_number
            certificate_number = row.certificate_number
            normalized_certificate = _normalize_certificate_number(certificate_number)
            if normalized_certificate in seen_certificates:
                result_rows.append(
                    EquipmentSIBulkImportRowRead(
                        row_number=row_number,
                        certificate_number=certificate_number,
                        status="skipped",
                        message="Duplicate certificate number inside the uploaded file.",
                    )
                )
                continue
            seen_certificates.add(normalized_certificate)

            try:
                search_results = await arshin.search_by_certificate(
                    certificate_number=certificate_number,
                    year=row.verification_year,
                )
                matched_result = _select_bulk_import_candidate(
                    certificate_number=certificate_number,
                    results=search_results,
                )
                if matched_result is None:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                        detail="Arshin search returned multiple ambiguous records.",
                    )

                detail = await arshin.get_vri_detail(vri_id=matched_result.vri_id)
                equipment = self.create_equipment(
                    _build_si_create_request_from_arshin(
                        folder_id=folder_id,
                        object_name=object_name,
                        status_value=status_value,
                        current_location_manual=current_location_manual,
                        result=matched_result,
                        detail=detail,
                    ),
                    current_user=None,
                )
                result_rows.append(
                    EquipmentSIBulkImportRowRead(
                        row_number=row_number,
                        certificate_number=certificate_number,
                        status="created",
                        message="SI equipment item was created from Arshin.",
                        equipment_id=equipment.id,
                        equipment_name=equipment.name,
                        vri_id=matched_result.vri_id,
                    )
                )
            except HTTPException as exc:
                row_status = "skipped" if exc.status_code == status.HTTP_409_CONFLICT else "error"
                result_rows.append(
                    EquipmentSIBulkImportRowRead(
                        row_number=row_number,
                        certificate_number=certificate_number,
                        status=row_status,
                        message=str(exc.detail),
                    )
                )

        result = EquipmentSIBulkImportResultRead(
            total_rows=len(rows),
            created_count=sum(1 for row in result_rows if row.status == "created"),
            skipped_count=sum(1 for row in result_rows if row.status == "skipped"),
            error_count=sum(1 for row in result_rows if row.status == "error"),
            rows=result_rows,
        )
        if current_user is not None:
            folder = self._get_folder(folder_id)
            self._record_event(
                category=EventCategory.EQUIPMENT,
                action="si_imported",
                user=current_user,
                title=f"Выполнен импорт СИ в папку «{folder.name}»",
                description=(
                    f"Создано: {result.created_count}. "
                    f"Пропущено: {result.skipped_count}. "
                    f"Ошибок: {result.error_count}."
                ),
                folder_id=folder.id,
                folder_name=folder.name,
            )
        return result

    def export_equipment_registry_xlsx(
        self,
        *,
        folder_id: int | None = None,
        group_id: int | None = None,
        equipment_ids: list[int] | None = None,
        query: str | None = None,
        status_value: EquipmentStatus | None = None,
        equipment_type: EquipmentType | None = None,
    ) -> Path:
        equipment_items = self.list_equipment(
            folder_id=folder_id,
            group_id=group_id,
            equipment_ids=equipment_ids,
            query=query,
            status=status_value,
            equipment_type=equipment_type,
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Equipment"
        headers = [
            "Папка",
            "Категория",
            "Статус",
            "Наименование",
            "Модификация",
            "Заводской номер",
            "Год выпуска",
            "Объект",
            "Локация",
            "Номер Аршина",
            "Действительно до",
        ]
        sheet.append(headers)

        for equipment in equipment_items:
            folder_name = None
            if equipment.folder_id is not None:
                folder = self.folders.get_by_id(equipment.folder_id)
                folder_name = folder.name if folder is not None else None

            verification = equipment.si_verification
            sheet.append(
                [
                    folder_name,
                    equipment.equipment_type.value,
                    equipment.status.value,
                    equipment.name,
                    equipment.modification,
                    equipment.serial_number,
                    equipment.manufacture_year,
                    equipment.object_name,
                    equipment.current_location_manual,
                    verification.result_docnum if verification is not None else None,
                    (
                        verification.valid_date.strftime("%d.%m.%Y")
                        if verification is not None and verification.valid_date is not None
                        else None
                    ),
                ]
            )

        return _save_workbook_to_temp_file(workbook)

    def _build_workbook_bytes(
        self,
        *,
        sheet_title: str,
        headers: list[str],
        rows: list[list[object | None]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_title
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        return _save_workbook_to_temp_file(workbook)

    def _format_sheet_date(self, value: date | datetime | None) -> str | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            value = value.date()
        return value.strftime("%d.%m.%Y")
