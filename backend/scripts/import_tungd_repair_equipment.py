from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import PROJECT_ROOT
from app.db.session import SessionLocal
from app.models.equipment import Equipment, EquipmentComment, EquipmentFolder, EquipmentType
from app.models.user import User
from app.schemas.equipment import EquipmentCommentCreateRequest, EquipmentCreateRequest
from app.services.equipment_service import EquipmentService
from app.services.user_service import UserService

DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "docs" / "13. ДЛЯ НАС_25.03.26.xlsx"
DEFAULT_SHEET_NAME = "Общая информация"
DEFAULT_FOLDER_NAME = "ТЮНГД Ремонт СИ"
HEADER_ROW_MAX = 5
DATA_START_ROW = 6
COMMENT_HEADER_MARKER = "КОММЕНТАРИ"


@dataclass(frozen=True, slots=True)
class ImportRow:
    row_number: int
    object_name: str
    current_location_manual: str | None
    name: str
    modification: str | None
    serial_number: str | None
    comment_text: str | None


@dataclass(slots=True)
class ImportSummary:
    parsed_rows: int = 0
    created_count: int = 0
    existing_count: int = 0
    comments_added_count: int = 0
    comments_skipped_count: int = 0
    error_count: int = 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Импортирует оборудование из Excel в папку metroLog как категорию OTHER "
            "и добавляет последний непустой комментарий."
        )
    )
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_WORKBOOK_PATH,
        help=f"Путь к Excel-файлу. По умолчанию: {DEFAULT_WORKBOOK_PATH}",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET_NAME,
        help=f"Имя листа. По умолчанию: {DEFAULT_SHEET_NAME}",
    )
    parser.add_argument(
        "--folder-name",
        default=DEFAULT_FOLDER_NAME,
        help=f"Имя целевой папки metroLog. По умолчанию: {DEFAULT_FOLDER_NAME}",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Ограничить количество обрабатываемых строк.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Применить изменения. Без флага выполняется только dry-run.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = args.xlsx.expanduser().resolve()
    rows = load_import_rows(
        workbook_path=workbook_path,
        sheet_name=args.sheet,
        fallback_object_name=args.folder_name,
    )
    if args.limit is not None:
        rows = rows[: args.limit]

    print(f"Источник: {workbook_path}")
    print(f"Лист: {args.sheet}")
    print(f"Целевая папка: {args.folder_name}")
    print(f"Найдено строк для импорта: {len(rows)}")
    print(f"Режим: {'apply' if args.apply else 'dry-run'}")

    summary = ImportSummary(parsed_rows=len(rows))
    with SessionLocal() as session:
        admin_user = UserService(session).ensure_bootstrap_admin()
        if admin_user is None:
            raise RuntimeError("Bootstrap admin is not available.")
        folder = session.scalar(
            select(EquipmentFolder).where(EquipmentFolder.name == args.folder_name)
        )
        if folder is None:
            raise RuntimeError(f"Folder '{args.folder_name}' not found.")

        service = EquipmentService(session, access_user=admin_user)
        for row in rows:
            try:
                process_row(
                    session=session,
                    service=service,
                    folder=folder,
                    admin_user=admin_user,
                    row=row,
                    apply_changes=args.apply,
                    summary=summary,
                )
            except Exception as exc:  # pragma: no cover - reporting path
                summary.error_count += 1
                if args.apply:
                    session.rollback()
                print(f"[ERROR] row={row.row_number}: {exc}")

    print_summary(summary=summary)
    return 0 if summary.error_count == 0 else 1


def load_import_rows(
    *,
    workbook_path: Path,
    sheet_name: str,
    fallback_object_name: str,
) -> list[ImportRow]:
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[sheet_name]
    comment_columns = find_comment_columns(sheet)
    rows: list[ImportRow] = []
    for row_number in range(DATA_START_ROW, sheet.max_row + 1):
        raw_name = normalize_text(sheet.cell(row=row_number, column=4).value)
        name = normalize_equipment_name(raw_name)
        modification = normalize_text(sheet.cell(row=row_number, column=5).value)
        serial_number = normalize_text(sheet.cell(row=row_number, column=6).value)
        if all(value is None for value in (name, modification, serial_number)):
            continue
        if name is None:
            raise RuntimeError(f"Equipment name is empty at row {row_number}.")

        technology_site = normalize_text(sheet.cell(row=row_number, column=2).value)
        dismantle_place = normalize_text(sheet.cell(row=row_number, column=3).value)
        object_name = technology_site or dismantle_place or fallback_object_name
        current_location_manual = (
            dismantle_place if dismantle_place and dismantle_place != object_name else None
        )
        comment_text = get_latest_comment_text(
            sheet=sheet,
            row_number=row_number,
            columns=comment_columns,
        )

        rows.append(
            ImportRow(
                row_number=row_number,
                object_name=object_name,
                current_location_manual=current_location_manual,
                name=name,
                modification=modification,
                serial_number=serial_number,
                comment_text=comment_text,
            )
        )
    return rows


def find_comment_columns(sheet) -> list[int]:
    columns: list[int] = []
    for column_number in range(1, sheet.max_column + 1):
        header_values = [
            normalize_text(sheet.cell(row=row_number, column=column_number).value)
            for row_number in range(1, HEADER_ROW_MAX + 1)
        ]
        if any(value and COMMENT_HEADER_MARKER in value.upper() for value in header_values):
            columns.append(column_number)
    return columns


def get_latest_comment_text(*, sheet, row_number: int, columns: list[int]) -> str | None:
    for column_number in reversed(columns):
        value = normalize_text(sheet.cell(row=row_number, column=column_number).value)
        if value is not None:
            return value
    return None


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r\n", "\n").strip()
    return text or None


def normalize_equipment_name(value: str | None) -> str | None:
    if value is None:
        return None
    if len(value) <= 255:
        return value
    short_candidate = value.split("(", 1)[0].strip()
    if short_candidate and len(short_candidate) <= 255:
        return short_candidate
    return value[:255].rstrip()


def process_row(
    *,
    session: Session,
    service: EquipmentService,
    folder: EquipmentFolder,
    admin_user: User,
    row: ImportRow,
    apply_changes: bool,
    summary: ImportSummary,
) -> None:
    existing = session.scalar(
        select(Equipment)
        .where(
            Equipment.folder_id == folder.id,
            Equipment.object_name == row.object_name,
            Equipment.name == row.name,
            Equipment.modification.is_(row.modification)
            if row.modification is None
            else Equipment.modification == row.modification,
            Equipment.serial_number.is_(row.serial_number)
            if row.serial_number is None
            else Equipment.serial_number == row.serial_number,
        )
        .order_by(Equipment.id.asc())
    )

    if existing is None:
        if not apply_changes:
            summary.created_count += 1
            if row.comment_text:
                summary.comments_added_count += 1
            print(
                f"[DRY-RUN][CREATE] row={row.row_number} "
                f"{row.name} / {row.modification or '—'} / {row.serial_number or '—'}"
            )
            return

        existing = service.create_equipment(
            EquipmentCreateRequest(
                folder_id=folder.id,
                object_name=row.object_name,
                equipment_type=EquipmentType.OTHER,
                name=row.name,
                modification=row.modification,
                serial_number=row.serial_number,
                status="IN_WORK",
                current_location_manual=row.current_location_manual,
            ),
            current_user=admin_user,
        )
        summary.created_count += 1
        print(
            f"[CREATED] row={row.row_number} equipment_id={existing.id} "
            f"{existing.name} / {existing.modification or '—'} / {existing.serial_number or '—'}"
        )
    else:
        summary.existing_count += 1
        print(
            f"[EXISTS] row={row.row_number} equipment_id={existing.id} "
            f"{existing.name} / {existing.modification or '—'} / {existing.serial_number or '—'}"
        )

    if row.comment_text is None:
        summary.comments_skipped_count += 1
        return

    comment_exists = session.scalar(
        select(EquipmentComment)
        .where(
            EquipmentComment.equipment_id == existing.id,
            EquipmentComment.text == row.comment_text,
            EquipmentComment.is_private.is_(False),
        )
        .limit(1)
    )
    if comment_exists is not None:
        summary.comments_skipped_count += 1
        return

    if not apply_changes:
        summary.comments_added_count += 1
        return

    service.create_comment(
        equipment_id=existing.id,
        payload=EquipmentCommentCreateRequest(text=row.comment_text),
        author=admin_user,
    )
    summary.comments_added_count += 1
    print(f"[COMMENT] row={row.row_number} equipment_id={existing.id}")


def print_summary(*, summary: ImportSummary) -> None:
    print("")
    print("Итог:")
    print(f"  parsed_rows={summary.parsed_rows}")
    print(f"  created_count={summary.created_count}")
    print(f"  existing_count={summary.existing_count}")
    print(f"  comments_added_count={summary.comments_added_count}")
    print(f"  comments_skipped_count={summary.comments_skipped_count}")
    print(f"  error_count={summary.error_count}")


if __name__ == "__main__":
    raise SystemExit(main())
